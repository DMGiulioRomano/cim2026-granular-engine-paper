#!/usr/bin/env python3
"""
make_pptx.py — versione PPTX di slides/index.html per Google Presentazioni
(issue #40: il CIM proietta con Google Presentazioni, reveal.js non è previsto).

Testi, YAML, note del relatore e riquadri del cursore (data-axes, data-t,
data-page) si leggono da index.html, che resta la fonte: qui si decide solo
l'impaginazione. Il cursore sincronizzato con l'audio non esiste in Google
Presentazioni, quindi ogni MAP diventa un video: MAP + cursore rosso + audio
dell'esempio, nella stessa logica dello script di index.html (pagine da 10 s
per complete_example, pannelli (a) e (b) in sequenza per deviation).

Prerequisito: `make slides` (media in slides/media/). Output in slides/media/:
    pptx/<nome>.png  .mp4      poster statico e video di ogni MAP
    pge_cim2026.pptx            il deck, video incorporati (PowerPoint/Keynote)

In Google Presentazioni i video incorporati non vengono importati: vanno
reinseriti da Drive (Inserisci > Video) sopra il poster, che resta al suo posto.

Uso:
    python slides/make_pptx.py
"""
import glob
import html
import os
import re
import subprocess
import textwrap

import qrcode
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SLIDES = os.path.join(REPO, "slides")
MEDIA = os.path.join(SLIDES, "media")
OUT = os.path.join(MEDIA, "pptx")
DECK = os.path.join(MEDIA, "pge_cim2026.pptx")

VIDEO_W = 1920          # larghezza dei video delle MAP, in px
CURSOR_W = 4            # spessore del cursore a VIDEO_W
FPS = 25
SANS = "Source Sans Pro"   # il font del tema white di reveal, in Google Fonts
MONO = "Source Code Pro"
INK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x55, 0x55, 0x55)
CODE_BG = RGBColor(0x27, 0x28, 0x22)   # monokai, come plugin/highlight
REPO_URL = "https://github.com/DMGiulioRomano/PythonGranularEngine"


# --- lettura di index.html -------------------------------------------------

def parse_index():
    src = open(os.path.join(SLIDES, "index.html"), encoding="utf-8").read()
    sections = []
    for sec in re.findall(r"<section.*?</section>", src, re.S):
        title = re.search(r"<h[23]>(.*?)</h[23]>", sec, re.S).group(1)
        codes = [textwrap.dedent(html.unescape(c)).strip("\n")
                 for c in re.findall(r"<code[^>]*data-trim[^>]*>(.*?)</code>", sec, re.S)]
        notes = re.search(r'<aside class="notes">(.*?)</aside>', sec, re.S)
        notes = " ".join(notes.group(1).split()) if notes else ""
        audios = re.findall(r'<audio[^>]*src="([^"]+)"', sec)
        heads = []
        for div in re.findall(r'<div class="playhead".*?</div>', sec, re.S):
            imgs = re.findall(r'<img class="(\w+)"[^>]*?(?:data-)?src="([^"]+)"', div)
            if not imgs or imgs[0][0] != "map":
                continue        # lo spettrogramma di distribution resta statico
            attr = lambda k: (re.search(rf'data-{k}="([^"]+)"', div) or [None, None])[1]
            heads.append({
                "imgs": [s for _, s in imgs],
                "axes": [list(map(float, a.split())) for a in attr("axes").split(";")],
                "t": list(map(float, attr("t").split())) if attr("t") else None,
                "page": float(attr("page")) if attr("page") else None,
            })
        sections.append({"title": title, "codes": codes, "notes": notes,
                         "audios": audios, "heads": heads, "html": sec})
    return sections


# --- media -----------------------------------------------------------------

def run(*cmd):
    subprocess.run(cmd, check=True)


def duration(path):
    out = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration",
                          "-of", "csv=p=0", path], check=True, capture_output=True, text=True)
    return float(out.stdout)


def raster(src):
    """PNG ad alta risoluzione di un'immagine di index.html (media/...)."""
    base = os.path.splitext(os.path.basename(src))[0]
    path = os.path.join(SLIDES, src)
    if src.endswith(".png"):
        return path
    png = os.path.join(OUT, base + ".png")
    if not os.path.exists(png):
        pdf = glob.glob(os.path.join(REPO, "paper", "examples", "*", base + ".pdf"))
        if pdf:     # dal PDF dell'esempio, più veloce delle svg da decine di MB
            run("pdftoppm", "-png", "-singlefile", "-scale-to-x", "2560", "-scale-to-y", "-1",
                pdf[0], png[:-4])
        else:
            run("rsvg-convert", "-w", "2560", "-o", png, path)
    return png


def segment(out, img, axes, span, audio, ss, dur, size):
    """Un tratto di video: immagine fissa, cursore da x0 a x1 in `span` secondi."""
    W, H = size
    x0, x1, y0, y1 = axes
    X0, DX = W * x0 / 100, W * (x1 - x0) / 100
    Y0, CH = round(H * y0 / 100), round(H * (y1 - y0) / 100)
    run("ffmpeg", "-loglevel", "error", "-y",
        "-loop", "1", "-framerate", str(FPS), "-i", img,
        "-ss", str(ss), "-t", str(dur), "-i", audio,
        "-f", "lavfi", "-i", f"color=c=0xdd0000:s={CURSOR_W}x{CH}:r={FPS}",
        "-filter_complex",
        f"[0]scale={W}:{H},format=yuv420p[bg];"
        f"[bg][2]overlay=x='{X0:.2f}+{DX:.2f}*t/{span}':y={Y0}:shortest=1[v]",
        "-map", "[v]", "-map", "1:a", "-t", str(dur),
        "-c:v", "libx264", "-tune", "stillimage", "-pix_fmt", "yuv420p", "-r", str(FPS),
        "-c:a", "aac", "-b:a", "192k", out)


def map_video(name, head, audios):
    """Video della MAP di una slide, con la stessa logica del cursore di index.html."""
    pngs = [raster(s) for s in head["imgs"]]
    w, h = Image.open(pngs[0]).size
    size = (VIDEO_W, 2 * round(VIDEO_W * h / w / 2))
    tracks = [os.path.join(SLIDES, a) for a in audios]
    segs = []   # (immagine, riquadro, span, audio, ss, dur)
    if head["page"]:            # pagine sovrapposte, un solo audio
        page, total = head["page"], duration(tracks[0])
        for i, (png, axes) in enumerate(zip(pngs, head["axes"])):
            segs.append((png, axes, page, tracks[0], i * page, min(page, total - i * page)))
    elif head["t"]:             # la MAP è una pagina: solo il tratto che copre
        t0, t1 = head["t"]
        segs.append((pngs[0], head["axes"][0], t1 - t0, tracks[0], t0, t1 - t0))
    else:                       # un riquadro per audio, in sequenza
        for axes, track in zip(head["axes"], tracks):
            d = duration(track)
            segs.append((pngs[0], axes, d, track, 0, d))
    parts = []
    for i, seg in enumerate(segs):
        parts.append(os.path.join(OUT, f"{name}_seg{i}.mp4"))
        segment(parts[-1], *seg, size)
    mp4 = os.path.join(OUT, name + ".mp4")
    if len(parts) == 1:
        os.replace(parts[0], mp4)
    else:
        lst = os.path.join(OUT, name + "_concat.txt")
        open(lst, "w").write("".join(f"file '{p}'\n" for p in parts))
        run("ffmpeg", "-loglevel", "error", "-y", "-f", "concat", "-safe", "0",
            "-i", lst, "-c", "copy", mp4)
        for p in parts + [lst]:
            os.remove(p)
    return pngs[0], mp4


def diagram_png(sec):
    svg = re.search(r"<svg.*?</svg>", sec["html"], re.S).group(0)
    svg = svg.replace("<svg ", '<svg xmlns="http://www.w3.org/2000/svg" ', 1)
    path = os.path.join(OUT, "diagram.svg")
    open(path, "w", encoding="utf-8").write(svg)
    png = path[:-4] + ".png"
    run("rsvg-convert", "-w", "2400", "-b", "white", "-o", png, path)
    return png


def qr_png():
    png = os.path.join(OUT, "qr_repo.png")
    qrcode.make(REPO_URL, border=1).save(png)
    return png


# --- impaginazione ---------------------------------------------------------
# Coordinate in pollici su 13,333 x 7,5: la slide di reveal (1280 x 720 px) a 96 dpi.

def text(slide, x, y, w, h, paras, size=18, align=PP_ALIGN.LEFT, font=SANS,
         color=INK, anchor=MSO_ANCHOR.TOP):
    """paras: lista di paragrafi; ciascuno è una stringa o una lista di run
    (testo, {"i":..., "b":..., "sub":..., "sup":..., "mono":...})."""
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        for t, fmt in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            r.font.size = Pt(fmt.get("size", size))
            r.font.name = MONO if fmt.get("mono") else font
            r.font.italic = fmt.get("i")
            r.font.bold = fmt.get("b")
            r.font.color.rgb = fmt.get("color", color)
            if fmt.get("sub") or fmt.get("sup"):
                r.font._rPr.set("baseline", "-25000" if fmt.get("sub") else "30000")
    return tf


def title(slide, t):
    text(slide, 0.5, 0.25, 12.333, 0.75, [t], size=36, align=PP_ALIGN.CENTER)
    slide.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True


YAML_COLORS = {"key": RGBColor(0xF9, 0x26, 0x72), "str": RGBColor(0xE6, 0xDB, 0x74),
               "num": RGBColor(0xAE, 0x81, 0xFF), "com": RGBColor(0x75, 0x71, 0x5E),
               "txt": RGBColor(0xF8, 0xF8, 0xF2)}
TOKEN = re.compile(r'(?P<com>#.*)|(?P<str>"[^"]*")|(?P<num>-?\d*\.?\d+)|(?P<txt>[^#"\d-]+|.)')


def yaml_runs(line):
    m = re.match(r"(\s*-?\s*)([\w]+)(:)(.*)", line)
    runs = []
    if m:
        lead, key, colon, line = m.groups()
        runs += [(lead, "txt"), (key, "key"), (colon, "txt")]
    runs += [(t.group(0), t.lastgroup) for t in TOKEN.finditer(line)]
    return [(t, {"mono": True, "color": YAML_COLORS[k]}) for t, k in runs]


def code(slide, x, y, w, src, max_size=14, max_h=6.0):
    """Blocco YAML su sfondo monokai; il corpo si adatta a larghezza e altezza.
    Ritorna la y del bordo inferiore."""
    lines = src.split("\n")
    pad = 0.12
    size = min(max_size,
               (w - 2 * pad) * 72 / (0.62 * max(map(len, lines))),
               (max_h - 2 * pad) * 72 / (1.2 * len(lines)))
    h = len(lines) * size * 1.2 / 72 + 2 * pad
    box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = CODE_BG
    box.line.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = False    # una riga di YAML non va mai a capo
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(pad)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for k, line in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = Pt(size * 1.2)
        for t, fmt in yaml_runs(line) or [(" ", {"mono": True})]:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.name = MONO
            r.font.color.rgb = fmt.get("color", YAML_COLORS["txt"])
    return y + h


def fit(png, x, y, w, h):
    """Box (x, y, w, h) dell'immagine adattata al riquadro, centrata in orizzontale."""
    iw, ih = Image.open(png).size
    s = min(w / iw, h / ih)
    return x + (w - iw * s) / 2, y, iw * s, ih * s


def picture(slide, png, x, y, w, h):
    bx, by, bw, bh = fit(png, x, y, w, h)
    slide.shapes.add_picture(png, Inches(bx), Inches(by), Inches(bw), Inches(bh))
    return bx, by, bw, bh


def movie(slide, mp4, poster, x, y, w, h):
    """Poster come immagine (resta dopo l'import in Google Presentazioni) e,
    sopra, il video incorporato con lo stesso poster."""
    bx, by, bw, bh = picture(slide, poster, x, y, w, h)
    slide.shapes.add_movie(mp4, Inches(bx), Inches(by), Inches(bw), Inches(bh),
                           poster_frame_image=poster, mime_type="video/mp4")
    return bx, by, bw, bh


def caption(slide, x, y, w, paras, size=12):
    text(slide, x, y, w, 0.3 * len(paras), paras, size=size, color=GREY)


def declare_notes_master(prs):
    """python-pptx crea il notes master e la sua relazione ma non lo dichiara
    in <p:notesMasterIdLst>: PowerPoint e LibreOffice lo tollerano, Keynote
    rifiuta il file («formato non valido»)."""
    from pptx.oxml.ns import qn
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    rid = next(r.rId for r in prs.part.rels.values() if r.reltype == RT.NOTES_MASTER)
    root = prs.part._element
    if root.find(qn("p:notesMasterIdLst")) is None:
        lst = root.makeelement(qn("p:notesMasterIdLst"), {})
        lst.append(root.makeelement(qn("p:notesMasterId"), {qn("r:id"): rid}))
        root.find(qn("p:sldIdLst")).addnext(lst)


I = {"i": True}
SUB = {"sub": True}
ISUB = {"i": True, "sub": True}
B = {"b": True, "i": True}


def build(sections):
    os.makedirs(OUT, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    vids = {}

    def new(sec, n):
        slide = prs.slides.add_slide(blank)
        if n > 1:
            text(slide, 12.2, 7.05, 0.9, 0.3, [str(n)], size=12, color=GREY, align=PP_ALIGN.RIGHT)
        slide.notes_slide.notes_text_frame.text = sec["notes"]
        return slide

    def mapvid(sec, key):
        if key not in vids:
            vids[key] = map_video(key, sec["heads"][0], sec["audios"])
        return vids[key]

    n = 0
    for sec in sections:
        n += 1
        t = sec["title"]
        if t == "PythonGranularEngine":
            s = new(sec, n)
            text(s, 0.8, 2.0, 11.733, 0.9, ["PythonGranularEngine"], size=48,
                 align=PP_ALIGN.CENTER)
            s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
            text(s, 0.8, 3.0, 11.733, 1.1, ["un ambiente dichiarativo per la granulazione",
                                            "e la sua mappa sinottica"],
                 size=28, align=PP_ALIGN.CENTER)
            text(s, 0.8, 4.6, 11.733, 0.8, ["Giulio Romano De Mattia — Conservatorio «A. Casella», L'Aquila",
                                            "XXV CIM 2026"], size=18, align=PP_ALIGN.CENTER)
            text(s, 0.8, 5.7, 11.733, 0.4, ["github.com/DMGiulioRomano/PythonGranularEngine · "
                                            "doi.org/10.5281/zenodo.22996019"],
                 size=16, align=PP_ALIGN.CENTER, color=GREY)

        elif t == "L'ambiente":
            s = new(sec, n)
            title(s, t)
            picture(s, diagram_png(sec), 0.5, 1.15, 12.333, 5.0)
            caption(s, 0.5, 6.45, 12.333, ["notazione prescrittiva e descrittiva nel senso di Seeger (1958)"],
                    size=16)
            s.shapes[-1].text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

        elif t.startswith("Composizione"):
            s = new(sec, n)
            title(s, t)
            make, sco = (os.path.join(MEDIA, f"complete_example_{k}.mp4") for k in ("make", "sco"))
            for mp4, y, h in ((make, 1.15, 2.25), (sco, 3.5, 1.575)):
                poster = os.path.join(OUT, os.path.basename(mp4)[:-4] + "_poster.png")
                run("ffmpeg", "-loglevel", "error", "-y", "-ss", "3", "-i", mp4,
                    "-frames:v", "1", poster)
                movie(s, mp4, poster, 0.5, y, 3.0, h)
            text(s, 0.5, 5.2, 3.0, 0.6, [[("38.072", {"b": True})]], size=40)
            text(s, 0.5, 5.85, 3.0, 1.0, ["grani in 32 secondi",
                                          [("complete_example.yml", {"mono": True, "size": 13})],
                                          "back-end Csound"], size=14)
            poster, mp4 = mapvid(sec, "complete_example_10-15s")
            movie(s, mp4, poster, 3.8, 1.15, 9.05, 6.2)

        elif t.startswith("Le condizioni"):
            s = new(sec, n)
            title(s, t)
            code(s, 0.6, 1.4, 3.6, sec["codes"][0], max_size=18)
            poster, mp4 = mapvid(sec, "identity")
            movie(s, mp4, poster, 4.6, 1.15, 8.25, 6.2)

        elif t.startswith("La posizione"):
            s = new(sec, n)
            title(s, t)
            code(s, 0.6, 1.4, 4.0, sec["codes"][0], max_size=18)
            poster, mp4 = mapvid(sec, "pointer")
            movie(s, mp4, poster, 4.9, 1.15, 7.95, 6.2)

        elif t == "Inter-onset time":
            s = new(sec, n)
            title(s, t)
            y = code(s, 0.5, 1.15, 4.3, sec["codes"][0], max_size=14)
            text(s, 0.5, y + 0.2, 4.3, 0.7, [
                [("IOT(", {}), ("d", I), (") = (1 − ", {}), ("d", I), (") IOT", {}), ("avg", SUB),
                 (" + ", {}), ("d", I), (" 𝒰(0, 2 IOT", {}), ("avg", SUB), ("),  ", {}),
                 ("d", I), (" ∈ [0, 1]", {})],
                [("d", I), (" = ", {}), ("distribution", {"mono": True}), (" · IOT", {}),
                 ("avg", SUB), (" = 1/", {}), ("density", {"mono": True})]],
                size=13, font="Times New Roman")
            spec = os.path.join(MEDIA, "distribution_spectrogram.png")
            picture(s, spec, 0.5, y + 1.05, 4.3, 7.2 - (y + 1.05))
            poster, mp4 = mapvid(sec, "distribution")
            movie(s, mp4, poster, 5.1, 1.15, 7.75, 6.2)

        elif t.startswith("Ampiezza"):
            s = new(sec, n)
            title(s, t)
            common, a, b = sec["codes"]
            y = code(s, 0.5, 1.1, 4.3, common, max_size=11)
            text(s, 0.5, y + 0.08, 4.3, 0.3, ["(a)"], size=14)
            y = code(s, 0.5, y + 0.38, 4.3, a, max_size=11)
            text(s, 0.5, y + 0.08, 4.3, 0.35, [[
                ("v", I), ("n", ISUB), (" = ", {}), ("c", I), ("(", {}), ("τ", I), ("n", ISUB),
                (") + ", {}), ("ξ", I), ("n", ISUB), (" ", {}), ("ρ", I), ("(", {}), ("τ", I),
                ("n", ISUB), ("),  ", {}), ("ξ", I), ("n", ISUB), (" ∼ 𝒰(−½, ½)", {})]],
                size=15, font="Times New Roman")
            text(s, 0.5, y + 0.5, 4.3, 0.3, ["(b)"], size=14)
            y = code(s, 0.5, y + 0.8, 4.3, b, max_size=11)
            text(s, 0.5, y + 0.08, 4.3, 0.35, [[
                ("gate  ", {"size": 12, "color": GREY}),
                ("v", I), ("n", ISUB), (" = ", {}), ("c", I), ("(", {}), ("τ", I), ("n", ISUB),
                (") + ", {}), ("g", B), ("n", {"b": True, "i": True, "sub": True}), (" ", {}),
                ("ξ", I), ("n", ISUB), (" ", {}), ("ρ", I), ("(", {}), ("τ", I), ("n", ISUB),
                ("),  ", {}), ("g", B), ("n", {"b": True, "i": True, "sub": True}),
                (" ∼ Bernoulli(", {}), ("p", I), ("(", {}), ("τ", I), ("n", ISUB), ("))", {})]],
                size=15, font="Times New Roman")
            poster, mp4 = mapvid(sec, "deviation")
            movie(s, mp4, poster, 5.1, 1.1, 7.75, 6.25)
            caption(s, 5.1, 7.05, 6.5, ["video: (a) poi (b), 12 s ciascuno"], size=11)

        elif t == "La sola probabilità":
            s = new(sec, n)
            title(s, t)
            y = code(s, 0.5, 1.15, 4.3, sec["codes"][0], max_size=13)
            rows = re.findall(r"<tr><td>(.*?)</td><td>(.*?)</td></tr>", sec["html"])
            tbl = s.shapes.add_table(len(rows) + 1, 2, Inches(0.5), Inches(y + 0.25),
                                     Inches(4.3), Inches(0.3 * (len(rows) + 1))).table
            for r, (k, v) in enumerate([("parametro", "jitter implicito")] + rows):
                for c, val in enumerate((k, v)):
                    cell = tbl.cell(r, c)
                    cell.text = html.unescape(re.sub(r"<sup>†</sup>", "†", val))
                    cell.margin_top = cell.margin_bottom = Inches(0.03)
                    font = cell.text_frame.paragraphs[0].runs[0].font
                    font.size, font.name, font.bold = Pt(13), SANS, r == 0
            caption(s, 0.5, y + 0.4 + 0.3 * (len(rows) + 1), 4.3,
                    ["† detune in ratio-space dopo la quantizzazione di griglia; "
                     "±0,0025 sul moltiplicatore per unità ratio."], size=10)
            poster, mp4 = mapvid(sec, "probability")
            movie(s, mp4, poster, 5.1, 1.15, 7.75, 6.2)

        elif t == "L'esempio completo":
            s = new(sec, n)
            title(s, t)
            code(s, 0.4, 1.05, 5.0, sec["codes"][0], max_size=10, max_h=6.3)
            poster, mp4 = mapvid(sec, "complete_example")
            movie(s, mp4, poster, 5.6, 1.05, 7.35, 6.3)
            caption(s, 5.6, 7.05, 6.5, ["video: la MAP in pagine da 10 s che seguono l'audio"], size=11)

        elif t.startswith("La molteplicità"):
            s = new(sec, n)
            title(s, t)
            code(s, 0.4, 1.05, 4.4, sec["codes"][0], max_size=10, max_h=6.1)
            poster, mp4 = mapvid(sec, "PGE_voices")
            movie(s, mp4, poster, 5.0, 1.05, 7.9, 6.3)

        elif t == "Conclusioni":
            s = new(sec, n)
            title(s, t)
            items = re.findall(r"<li>(.*?)</li>", sec["html"])
            tf = text(s, 1.0, 1.6, 7.3, 2.4, items, size=24)
            for p in tf.paragraphs:
                p.space_after = Pt(14)
                p.runs[0].text = "•  " + p.runs[0].text
            text(s, 1.0, 4.3, 7.3, 1.4, [
                [("Limiti", {"b": True})],
                "posizione, non contenuto · assi senza covarianza · altezza solo nel colore"],
                size=20, color=GREY)
            picture(s, qr_png(), 9.2, 1.6, 3.0, 3.0)
            caption(s, 8.7, 4.75, 4.0, ["github.com/DMGiulioRomano/PythonGranularEngine",
                                        "doi.org/10.5281/zenodo.22996019"], size=13)
            for p in s.shapes[-1].text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
        else:
            raise SystemExit(f"slide senza impaginazione: {t!r}")

    declare_notes_master(prs)
    prs.save(DECK)
    print(f"=== {DECK} ({n} slide) ===")


if __name__ == "__main__":
    build(parse_index())
